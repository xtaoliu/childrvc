"""Build R15 full draft from the author's R14."""
import copy, sys
sys.path.insert(0, "/home/user/childrvc/analysis/draft")
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docxhelp import replace_markup, find, set_text, insert_after, markup_runs, clear
from edits_A import A
from text_exp1 import EXP1_RESULTS, FIG_CAPTIONS
from text_exp2 import EXP2_RESULTS, TABLE2_CAPTION, TABLE2_HEADER, TABLE2_NOTE
from text_disc import DISC_SUMMARY, DISC, CONCLUSION, ABSTRACT, KEYWORDS
from text_sens import SENS

SRC = "/home/user/childrvc/drafts/Manuscript_R14_final_from_user.docx"
OUT = "/home/user/childrvc/drafts/Manuscript_R15_full_draft.docx"
FIG = "/home/user/childrvc/analysis/fig/"
FIGFILE = {"fig1": "fig1_stills.png", "fig2": "fig2_comprehension.png", "fig3": "fig3_association.png", "fig4": "fig4_items.png", "fig5": "fig5_production.png"}
doc = Document(SRC)

def delete(p):
    p._p.getparent().remove(p._p)

# ---------- Part A: participants ----------
for prefix, old, new in A:
    replace_markup(find(doc, prefix), old, new)

# ---------- Table 1: items actually used in the experiment ----------
t1 = doc.tables[0]
def set_cell(cell, it, rest=None):
    runs = [r for p in cell.paragraphs for r in p.runs]
    runs[0].text = it
    if rest is not None: runs[1].text = rest
def set_plain(cell, text):
    runs = [r for p in cell.paragraphs for r in p.runs]
    runs[0].text = text
    for r in runs[1:]: r.text = ""
for ri in (7, 8, 19, 20):
    set_cell(t1.rows[ri].cells[2], "kāi", "\n“open”")
set_cell(t1.rows[7].cells[4], "sī-kāi"); set_plain(t1.rows[7].cells[7], "2,810")
set_cell(t1.rows[8].cells[3], "chě", "\n“pull”"); set_cell(t1.rows[8].cells[4], "chě-kāi"); set_plain(t1.rows[8].cells[7], "840ᵈ")
set_cell(t1.rows[19].cells[4], "jiā-kāi"); set_plain(t1.rows[19].cells[7], "21")
set_cell(t1.rows[20].cells[4], "chuí-kāi"); set_plain(t1.rows[20].cells[7], "1")
note1 = find(doc, "Note. Likelihood was rated")
replace_markup(note1, "ᶜ Includes many abstract uses.", "ᶜ Includes many abstract uses. ᵈ Includes idiomatic uses such as *chě-kāi sǎngzi* \"pull-open throat,\" \"raise one's voice.\"")
p63 = find(doc, "We also examined how conventional each RVC is.")
replace_markup(p63, "(e.g., sī-pò liǎn “tear-torn face,” “fall out with someone”)", "(e.g., *chě-kāi sǎngzi* \"pull-open throat,\" \"raise one's voice\")")

# ---------- Abstract and keywords after the title ----------
title = doc.paragraphs[0]
h_abs = insert_after(title, "Abstract", style="Heading 1")
p_abs = insert_after(h_abs, ABSTRACT, style="Normal"); p_abs.paragraph_format.first_line_indent = Inches(0)
p_kw = insert_after(p_abs, "*Keywords:* " + KEYWORDS, style="Normal"); p_kw.paragraph_format.first_line_indent = Inches(0.5)
p_kw.paragraph_format.page_break_before = False
# introduction starts on a new page with the title repeated (APA)
p_t2 = insert_after(p_kw, title.text, style="Title"); p_t2.paragraph_format.page_break_before = True

# ---------- helpers for figures and tables ----------
def add_figure(after, key):
    num, ttl, note = FIG_CAPTIONS[key]
    a = insert_after(after, num, style="Table Number")
    b = insert_after(a, "*" + ttl + "*", style="Table Title")
    c = insert_after(b, "", style="Normal"); clear(c)
    c.paragraph_format.first_line_indent = Inches(0); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.keep_with_next = True
    c.add_run().add_picture(FIG + FIGFILE[key], width=Inches(6.0))
    d = insert_after(c, "*Note.* " + note, style="Table Note")
    return d

def add_table2(after):
    num, ttl = TABLE2_CAPTION
    a = insert_after(after, num, style="Table Number")
    b = insert_after(a, "*" + ttl + "*", style="Table Title")
    prod = pd.read_csv("/home/user/childrvc/analysis/prod.csv")
    ct = pd.crosstab([prod.age, prod.result], prod.category, normalize="index") * 100
    n = prod.groupby(["age", "result"]).size()
    rows = []
    for age, lab in (("4yo", "4-year-olds"), ("5yo", "5-year-olds"), ("adult", "Adults")):
        for res in ("full", "zero"):
            r = ct.loc[(age, res)]
            rows.append([lab if res == "full" else "", res.capitalize(), str(n[(age, res)])] +
                        [f"{r.get(c, 0):.0f}" for c in ("RVC", "separate", "result_only", "action_only", "neither")])
    tbl = doc.add_table(rows=1 + len(rows), cols=len(TABLE2_HEADER))
    tblPr = copy.deepcopy(t1._tbl.tblPr)
    for e in tblPr.findall(qn("w:tblLayout")): tblPr.remove(e)
    tbl._tbl.remove(tbl._tbl.tblPr); tbl._tbl.insert(0, tblPr)
    for i, row in enumerate([TABLE2_HEADER] + rows):
        for j, txt in enumerate(row):
            cell = tbl.rows[i].cells[j]; p = cell.paragraphs[0]; p.style = doc.styles["Table Text"]
            for e in markup_runs(p.add_run("")._r, txt): p._p.append(e)
            if j >= 2: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i == 0:
                tcPr = cell._tc.get_or_add_tcPr(); bd = tcPr.makeelement(qn("w:tcBorders"), {})
                bt = bd.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "4", qn("w:space"): "0", qn("w:color"): "000000"})
                bd.append(bt); tcPr.append(bd)
    b._p.addnext(tbl._tbl)
    note_p = copy.deepcopy(note1._p)
    tbl._tbl.addnext(note_p); np_ = Paragraph(note_p, b._parent); set_text(np_, TABLE2_NOTE)
    return np_

def add_tableA1(after):
    a = insert_after(after, "Table A1", style="Table Number")
    b = insert_after(a, "*Proportion of Descriptions Containing the Intended V1 in the Verb–Video Correspondence Norming*", style="Table Title")
    nv = pd.read_csv("/home/user/childrvc/analysis/norming_video.csv")
    py = {"bo": "bāo", "xian": "xiān", "bai": "bāi", "si": "sī", "chui": "chuī", "ti": "tī", "qia": "qiā", "tu": "tú", "ca": "cā", "jia": "jiā",
          "nie": "niē", "pai": "pāi", "an": "àn", "che": "chě", "hui": "huī", "tui": "tuī", "jiao": "jiào", "nao": "náo", "chong": "chōng"}
    hdr = ["Pair", "Stronger-association V1", "First", "All", "Weaker-association V1", "First", "All"]
    rows = []
    for k in range(1, 11):
        s_ = nv[(nv.pair == k) & (nv.assoc == "stronger")].iloc[0]; w_ = nv[(nv.pair == k) & (nv.assoc == "weaker")].iloc[0]
        wv = "chuí" if (k == 10) else py[w_.v1]
        rows.append([str(k) + ("\u1D43" if k in (7, 8) else ""), "*" + py[s_.v1] + "*", f"{s_['first']:.2f}", f"{s_['all']:.2f}", "*" + wv + "*", f"{w_['first']:.2f}", f"{w_['all']:.2f}"])
    tbl = doc.add_table(rows=1 + len(rows), cols=len(hdr))
    tblPr = copy.deepcopy(t1._tbl.tblPr)
    for e in tblPr.findall(qn("w:tblLayout")): tblPr.remove(e)
    tbl._tbl.remove(tbl._tbl.tblPr); tbl._tbl.insert(0, tblPr)
    for i, row in enumerate([hdr] + rows):
        for j, txt in enumerate(row):
            cell = tbl.rows[i].cells[j]; p = cell.paragraphs[0]; p.style = doc.styles["Table Text"]
            for e in markup_runs(p.add_run("")._r, txt): p._p.append(e)
            if j in (2, 3, 5, 6): p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i == 0:
                tcPr = cell._tc.get_or_add_tcPr(); bd = tcPr.makeelement(qn("w:tcBorders"), {})
                bd.append(bd.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "4", qn("w:space"): "0", qn("w:color"): "000000"})); tcPr.append(bd)
    b._p.addnext(tbl._tbl)
    note_p = copy.deepcopy(note1._p); tbl._tbl.addnext(note_p); np_ = Paragraph(note_p, b._parent)
    set_text(np_, "*Note.* First = proportion of participants whose first description contained the intended V1; All = proportion of all descriptions, including those given on request, that contained the intended V1 (*n* = 8 participants per video). The All values are used as the covariate in the sensitivity analysis of Experiment 1. ᵃ Excluded from the analyses of Experiment 1.")
    return np_

def write_block(anchor, items):
    """Insert (kind, text) items after anchor; return last paragraph."""
    cur = anchor
    for kind, text in items:
        if kind == "P": cur = insert_after(cur, text, style="Normal")
        elif kind == "H3": cur = insert_after(cur, text, style="Heading 3")
        elif kind == "H2": cur = insert_after(cur, text, style="Heading 2")
        elif kind == "FIG": cur = add_figure(cur, text)
        elif kind == "TABLE2": cur = add_table2(cur)
        elif kind == "SENS": cur = write_block(cur, SENS)
    return cur

# ---------- Experiment 1 ----------
p90 = find(doc, "Accuracy on filler trials was 100%")
p91 = find(doc, "[待补：按八对、两组方案重新分析后报告。")
set_text(p90, EXP1_RESULTS[0][1]); delete(p91)
write_block(p90, EXP1_RESULTS[1:])
p54 = find(doc, "Because each expression was paired with its own video")
replace_markup(p54, " [逐项数值放附录]", " Item-level values are given in Table A1 in the Appendix.")
replace_markup(p54, "appeared in 62% of first descriptions", "appeared in 74% of first descriptions")
p56 = find(doc, "Twenty native speakers of Mandarin (21–35 years")
replace_markup(p56, "the paper ended up torn?", "the paper ended up torn apart?")
replace_markup(p56, "wants to break the walnut.", "wants to crack the walnut open.")
p52 = find(doc, "Each of the 20 actions was filmed in a full-result and a zero-result version")
replace_markup(p52, " [Figure 1：full-result 与 zero-result 视频的截图]", " Figure 1 shows stills from the four videos of one pair.")
add_figure(p52, "fig1")
p80 = find(doc, "Children were tested individually in a quiet room in their kindergarten.")
replace_markup(p80, "Children were tested individually in a quiet room in their kindergarten. After each video, they heard the question and answered orally,",
  "Children and adults received the same instructions. They were told that they would watch short videos of everyday activities, that after each video the experimenter would say what they had seen, and that they should judge whether this was right or wrong. On the practice trial, which used a filler video, the experimenter said, *Wǒ kàn-dào zài zhège shìpín lǐ, zhège nǚhái pā-xià le* \"I saw in this video that this girl lay down,\" asked whether this was right or wrong, and then asked, *Zhège nǚhái pā-xià le ma?* \"Did this girl lie down?\" Participants answered in their own words, typically by repeating or negating the verb (e.g., *pā-xià le* \"lay down\"; *méi pā-xià* \"did not lie down\"), and their answers were coded as acceptance or rejection. Children were tested individually in a quiet room in their kindergarten. After each video, they heard the test question and answered orally,")
replace_markup(p80, "they received oral instructions from an experimenter and selected their answers themselves.", "they received the instructions orally from an experimenter and selected their answers themselves.")
p97 = find(doc, "The 40 critical videos were divided into four lists.")
replace_markup(p97, "watched each video, and were asked to describe what happened. [指导语原文]",
  "watched each video, and described it. Children were told that they would play a question-and-answer game. Before the test trials, they were familiarized with three of the objects (a candle, a lychee, and a walnut). After each video, the experimenter asked, *Nǐ kàn-dào le shénme ya?* \"What did you see?\" If a child responded with a verb, no further question was asked. If a child named only an object, the experimenter asked, *Nà zhège rén zài zuò shénme ya?* \"Then what is this person doing?\" If a child said that they did not know, the video was played again and the first question repeated; trials on which the child still did not know were not analyzed. Adults were told that they would watch short videos of everyday activities and should describe in one sentence what they saw in each video. The experimenter demonstrated the task with a filler video (*Yí ge nǚhái pā-xià le* \"A girl lay down\").")

# ---------- Experiment 2 method ----------
p100 = find(doc, "The coding unit is the complete response to a video")
replace_markup(p100, "Nonresponses and responses without a verb are excluded and reported.", "Trials without a recorded response are excluded and reported. In the present draft, responses were coded with item-specific lists of V1s, V2s, and equivalent result expressions, together with rules for negation and attempt verbs (e.g., *shìtú* \"try\"). ⟦规则编码初稿：需人工逐条核对，并由第二编码者独立编码 20% 的回答后报告一致性⟧")
p102 = find(doc, "Information. Each response is also coded")
replace_markup(p102, " [LA稿件把字句中的谓语一律编为 RVC，这里改了]", "")
p105 = find(doc, "Two binary outcomes are modeled")
replace_markup(p105, "and modeled if cell sizes permit. [待定：多分类模型还是仅描述]", "descriptively, because several cells contain few responses.")
replace_markup(p105, "Responses that affirm the target result in zero-result events are counted by age group.", "Responses that affirm the target result in zero-result events are counted by age group. The two models are also refitted with the eight retained pairs.")
p107 = find(doc, "[待补：按新的编码方案重新编码并分析后报告。]")
first = EXP2_RESULTS[0]; set_text(p107, first[1])
write_block(p107, EXP2_RESULTS[1:])

# ---------- General Discussion ----------
p109 = find(doc, "[本节为初稿")
set_text(p109, DISC_SUMMARY)
# move the question paragraph before the summary
p110 = find(doc, "The study asks three questions")
p109._p.addprevious(p110._p)
def replace_par(prefix, keys):
    p = find(doc, prefix); set_text(p, DISC[keys[0]]); cur = p
    for k in keys[1:]: cur = insert_after(cur, DISC[k], style="Normal")
replace_par("A larger effect of result realization on RVC", ["rvc", "rvc2", "rvc3"])
replace_par("If children reject V1 descriptions of zero-result events more often", ["assoc"])
replace_par("Why action–result association might matter", ["assoc_why"])
replace_par("If children and adults show similar association effects", ["assoc_age"])
replace_par("Because the stronger- and weaker-association expressions differ in both", ["assoc_conf"])
replace_par("Comparing RVC use with mention of the target result", ["prod", "prod2", "prod_instr", "prod3"])
replace_par("Several limitations follow from the design.", ["lim"])
set_text(find(doc, "[结论待根据结果改写]"), CONCLUSION)

# ---------- References ----------
delete(find(doc, "Skordos, D., & Papafragou, A. (2016)"))
sp = find(doc, "Speer, R. (2022)")
newref = copy.deepcopy(sp._p); sp._p.addprevious(newref)
set_text(Paragraph(newref, sp._parent), "Song, Z. T., & Haryu, E. (2025, August). *\"The end matters\": Endstate salience in change-of-state events affects Mandarin-speaking children's acceptance of verbal packaging patterns of action and result* [Poster presentation]. European Conference on Developmental Psychology (ECDP 2025), Prague, Czech Republic.")
newref2 = copy.deepcopy(sp._p); sp._p.addprevious(newref2)
set_text(Paragraph(newref2, sp._parent), "Song, Z. T., & Haryu, E. (2026). *Mandarin-speaking children can correctly produce change-of-state verbs: Limited influence of event state structure* [Poster presentation]. APS Global Psychological Science Summit.")

lastref = [p for p in doc.paragraphs if p.style.name == "Reference"][-1]
h = insert_after(lastref, "Appendix", style="Heading 1"); h.paragraph_format.page_break_before = True
add_tableA1(h)
doc.save(OUT); print("saved", OUT)
