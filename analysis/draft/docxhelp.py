import re, copy
from docx.enum.text import WD_COLOR_INDEX
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

TOKEN = re.compile(r"(\*[^*]+\*|⟦[^⟧]+⟧|\^[^^]+\^)")

def smart(text):
    text = re.sub(r'"([^"]*)"', "\u201C\\1\u201D", text)
    text = re.sub(r"(\w)'", "\\1\u2019", text)
    text = re.sub(r"(^|[\s(])'", "\\1\u2018", text)
    return text

def _new_run_like(run_el, text, italic=None, highlight=False, sup=False):
    r = copy.deepcopy(run_el)
    for t in r.findall(qn("w:t")): r.remove(t)
    rpr = r.find(qn("w:rPr"))
    if rpr is not None:
        for tag in ("w:i", "w:iCs", "w:highlight", "w:vertAlign"):
            for e in rpr.findall(qn(tag)): rpr.remove(e)
    else:
        rpr = r.makeelement(qn("w:rPr"), {}); r.insert(0, rpr)
    if italic:
        rpr.append(rpr.makeelement(qn("w:i"), {}))
    if highlight:
        h = rpr.makeelement(qn("w:highlight"), {qn("w:val"): "yellow"}); rpr.append(h)
    if sup:
        v = rpr.makeelement(qn("w:vertAlign"), {qn("w:val"): "superscript"}); rpr.append(v)
    t = r.makeelement(qn("w:t"), {}); t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); r.append(t)
    return r

def markup_runs(run_el, text):
    """Return list of run elements cloned from run_el for marked-up text."""
    out = []
    for seg in TOKEN.split(smart(text)):
        if not seg: continue
        if seg.startswith("*") and seg.endswith("*") and len(seg) > 1:
            out.append(_new_run_like(run_el, seg[1:-1], italic=True))
        elif seg.startswith("⟦"):
            out.append(_new_run_like(run_el, "[" + seg[1:-1] + "]", highlight=True))
        elif seg.startswith("^") and seg.endswith("^") and len(seg) > 1:
            out.append(_new_run_like(run_el, seg[1:-1], sup=True))
        else:
            out.append(_new_run_like(run_el, seg))
    return out

def replace_markup(p, old, new):
    """Replace `old` (may span non-italic runs) with marked-up `new`."""
    runs = p.runs; texts = [r.text for r in runs]; full = "".join(texts)
    k = full.find(old)
    assert k >= 0, ("not found", old[:60], p.text[:80])
    end = k + len(old); pos = 0; first = None
    for r, t in zip(runs, texts):
        a, b = pos, pos + len(t); pos = b
        if b <= k or a >= end: continue
        lo, hi = max(k, a) - a, min(end, b) - a
        if first is None:
            first = r
            before, after = t[:lo], t[hi:]
            r.text = before
            new_els = markup_runs(r._r, new)
            anchor = r._r
            for e in new_els:
                anchor.addnext(e); anchor = e
            if after:
                tail = _new_run_like(r._r, after, italic=bool(r.italic))
                anchor.addnext(tail)
        else:
            r.text = t[:lo] + t[hi:]

def find(doc, prefix):
    hits = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(hits) == 1, (prefix, len(hits)); return hits[0]

def clear(p):
    for el in list(p._p):
        if el.tag in (qn("w:r"), qn("w:hyperlink")): p._p.remove(el)

def set_text(p, text, template_run=None):
    tmpl = template_run if template_run is not None else (p.runs[0]._r if p.runs else None)
    if tmpl is None:
        p.add_run(""); tmpl = p.runs[0]._r
    tmpl = copy.deepcopy(tmpl)
    clear(p)
    for e in markup_runs(tmpl, text): p._p.append(e)

def insert_after(p, text, style=None):
    new = copy.deepcopy(p._p)
    tmpl = None
    for el in list(new):
        if el.tag == qn("w:r") and tmpl is None: tmpl = copy.deepcopy(el)
        if el.tag != qn("w:pPr"): new.remove(el)
    p._p.addnext(new)
    np_ = Paragraph(new, p._parent)
    if style is not None: np_.style = style
    if tmpl is None:
        np_.add_run(""); tmpl = np_.runs[0]._r; tmpl = copy.deepcopy(tmpl); clear(np_)
    for e in markup_runs(tmpl, text): np_._p.append(e)
    return np_
