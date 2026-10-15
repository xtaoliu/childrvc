"""R18: move method details, model specifications, coding manual, and Table 2 to the Supplemental Material.
Builds on the R17 main text and supplement (run build_r17.py first)."""
import copy, sys
sys.path.insert(0, "/home/user/childrvc/analysis/draft")
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docxhelp import replace_markup, find, set_text, insert_after, markup_runs, clear
import text_supp as T

D = "/home/user/childrvc/drafts/"
main = Document(D + "Manuscript_R17.docx")
old = Document(D + "Manuscript_R17_Supplemental_Material.docx")

def delete(p): p._p.getparent().remove(p._p)
def rep_all(doc, a, b):
    n = 0
    for p in doc.paragraphs:
        if a in p.text: replace_markup(p, a, b); n += 1
    return n

# ---------- 1. remap old supplement references (S2 -> S8, S3 -> S9) ----------
for a, b in [("Supplemental Material, Sections S2 and S3", "Supplemental Material, Sections S8 and S9"),
             ("Section S3 of the Supplemental Material", "Section S9 of the Supplemental Material"),
             ("(Supplemental Material, Section S3)", "(Supplemental Material, Section S9)"),
             ("(Supplemental Material, Section S2)", "(Supplemental Material, Section S8)"),
             ("Supplemental Material (Section S2)", "Supplemental Material (Section S8)"),
             ("Table S1 in the Supplemental Material", "Table S1 (Supplemental Material, Section S2)")]:
    rep_all(main, a, b)

# ---------- 2. condense main-text paragraphs ----------
set_text(find(main, "When the materials were designed, the V1s were classified"),
  "When the materials were designed, the V1s were classified as Strong and Weak V1s following Martin et al. (2020, 2026), on the basis of whether they are compatible with a preverbal time-span adverbial. We no longer use this classification, because the diagnostic is binary and characterizes verbs rather than the relation between an action and a particular result (Supplemental Material, Section S1). The present study instead characterizes the pairs by adult ratings of each action's association with its target result, described below. The labels stronger-association and weaker-association refer to these ratings, not to the original classification.")
set_text(find(main, "Twenty native speakers of Mandarin (21–35 years"),
  "Twenty native speakers of Mandarin (21–35 years, *M* = 28.2; 9 female) rated each of the 20 V1–object expressions in an online questionnaire with two blocks. In the likelihood block, each item presented a perfective sentence containing the expression (e.g., *Xiǎo Lǐ sī-le zhǐ* \"Xiao Li tore the paper\"), and raters judged how likely it was that the target result eventually occurred, from 0 (*cannot happen*) to 100 (*will certainly happen*). In the typicality block, they rated how typical the action was as a way of achieving the target result, from 1 (*very atypical*) to 7 (*very typical*). The likelihood block always preceded the typicality block. The instructions and the construction of the four lists are described in the Supplemental Material (Section S3). Because the ratings were collected for sentences naming the expressions, they characterize the action–result relation that each expression describes, not the particular videos. The ratings were collected after the child data had been collected and initially analyzed with the original Strong/Weak coding; the raters did not take part in the child experiments and received no information about children's responses.")
delete(find(main, "Four lists were constructed, and five raters completed each list."))
set_text(find(main, "For each rater and pair, we computed the difference"),
  "Across the 10 pairs, the expected stronger-association expressions received higher likelihood ratings (mean difference = 33.06 points, 95% CI [21.74, 44.70], *p* = .003) and higher typicality ratings (mean difference = 2.45 points, 95% CI [1.25, 3.59], *p* = .007), and the pair-level differences in the two ratings were correlated, *r* = .79, *p* = .005 (statistical procedure in Section S3). Mean ratings for each expression are shown in Table 1.")
set_text(find(main, "A pair was retained for the analyses of Experiment 1"),
  "A pair was retained for the analyses of Experiment 1 if the expected stronger-association expression received reliably higher ratings than its partner on both scales (*p* < .05, uncorrected). Eight pairs met this criterion. Two did not: In Pair 7 (*qiā* \"pinch\" vs. *jiào* \"call\" a person), the typicality difference was reversed, and in Pair 8 (*tú* \"paint\" vs. *náo* \"scratch\" a palm), neither difference was reliable (Section S3). Because action–result association is defined by the two ratings jointly, the actions in these pairs cannot be classified as stronger- or weaker-association actions. The two pairs remained in the experimental lists, so all participants saw all 10 pairs, but their trials are excluded from the analyses of Experiment 1. The original design included all 10 pairs, and an initial analysis of the child data used all of them with the original Strong/Weak coding. The adult ratings were collected, and Pairs 7 and 8 excluded, after that analysis. The exclusion criterion refers only to the rating data, not to children's responses, but because both steps followed an inspection of the child data, analyses with all 10 pairs are reported alongside the main analyses of Experiment 1 (Section S8).")
set_text(find(main, "We also examined how conventional each RVC is."),
  "We also examined how conventional each RVC is. As noted in the Introduction, whether a V1 and a V2 combine depends on how the action named by V1 relates to the result named by V2. If the two actions in a pair differ in their association with the target result, this difference should also be reflected in the language, such that the RVC formed with the stronger-association V1 is the more conventional one. Corpus frequency thus provides a check on the contrast between the paired V1s that does not rely on the adult ratings, although conventionality concerns the RVC as an expression rather than the action–result relation itself. We therefore searched the modern Chinese section of the CCL corpus (Zhan et al., 2019) for each of the 20 RVCs and counted the documents in which each RVC occurs in its resultative use, including figurative and idiomatic uses, which also indicate that the combination is conventional. The search and checking procedure is described in the Supplemental Material (Section S4). Document counts are shown in Table 1.")
replace_markup(find(main, "Note. Likelihood was rated"), "(see text)", "(see Supplemental Material, Section S4)")
replace_markup(find(main, "Form (RVC vs. V1) and Action–Result Association (Stronger vs. Weaker) were manipulated"),
  "to balance “yes” and “no” answers.", "to balance “yes” and “no” answers. The test sentences and filler questions are listed in the Supplemental Material (Section S5).")
p80 = find(main, "Children and adults received the same instructions.")
replace_markup(p80, p80.text[:p80.text.index("Children were tested individually")],
  "Children and adults received the same instructions: The experimenter would say what they had seen in each video, and they should judge whether this was right or wrong (Supplemental Material, Section S6). Participants answered in their own words, typically by repeating or negating the verb (e.g., *pā-xià le* \"lay down\"; *méi pā-xià* \"did not lie down\"), and their answers were coded as acceptance or rejection. ")
p82 = find(main, "Trial-level acceptance is modeled with Bayesian")
replace_markup(p82, "Factors are sum-coded. Fixed effects have Normal(0, 1.5) priors, standard deviations Exponential(1) priors, and correlations LKJ(2) priors. Convergence is assessed by R̂ < 1.01 and adequate effective sample sizes.",
  "Factors are sum-coded; priors, sampling settings, convergence criteria, random-effects structures, and full parameter estimates are reported in the Supplemental Material (Section S7).")
replace_markup(find(main, "The first model addresses RQ1."),
  " Random effects are intercepts for participants, items (a V1–object expression and its RVC), and pairs; a by-participant slope for Result Realization; and by-item slopes for Form, Result Realization, and Age Group.", "")
replace_markup(find(main, "The second model addresses RQ2"),
  "Random effects are intercepts for participants, items, and pairs; a by-participant slope for Result Realization; by-item slopes for Result Realization and Age Group; and a by-pair slope for Action–Result Association, which allows the difference between the two V1s of a pair to vary across pairs.",
  "The random effects include a by-pair slope for Action–Result Association, which allows the difference between the two V1s of a pair to vary across pairs.")
p97 = find(main, "The 40 critical videos were divided into four lists.")
t97 = p97.text; a = t97.index("Children were told that they would play"); b = t97.index("Children responded orally")
replace_markup(p97, t97[a:b],
  "After each video, children were asked, *Nǐ kàn-dào le shénme ya?* \"What did you see?\" If they named only an object, they were asked, *Nà zhège rén zài zuò shénme ya?* \"Then what is this person doing?\" Adults were asked to describe in one sentence what they saw in each video. The full instructions are given in the Supplemental Material (Section S6). ")
replace_markup(find(main, "The coding unit is the complete response"),
  "Trials without a recorded response are excluded and reported.", "Trials without a recorded response are excluded and reported. The coding criteria are given in the Supplemental Material (Section S10).")
p102 = find(main, "Information. Each response is also coded")
replace_markup(p102, p102.text[len("Information. "):],
  "Each response is also coded for whether it mentions the action and whether it mentions the target result, and whether a mentioned target result is affirmed, denied, or presented as the actor's goal, so that responses affirming the target result in a zero-result event can be identified. Only explicitly expressed information is coded.")
p105 = find(main, "Two binary outcomes are modeled")
replace_markup(p105, "with random intercepts for participants, items (the videos of one action), and pairs; a by-participant slope for Result Realization; and by-item slopes for Result Realization and Age Group.", "with the random-effects structure given in Section S7.")
replace_markup(p105, "descriptively, because several cells contain few responses.", "descriptively in the Supplemental Material (Section S11), because several cells contain few responses.")
replace_markup(p105, "The two models are also refitted with the eight retained pairs.", "The two models are also refitted with the eight retained pairs (Section S11).")
replace_markup(find(main, "Ten child trials (seven from 4-year-olds"), ", and Table 2 shows the distribution of response categories.", ".")

# ---------- 3. move Experiment 2 response categories and Table 2 ----------
h_rc = [p for p in main.paragraphs if p.style.name == "Heading 3" and p.text == "Response Categories"][0]
body = main.element.body
els = []; el = h_rc._p.getnext()
while True:
    els.append(el)
    if el.tag == qn("w:p") and Paragraph(el, None).text.startswith("Restricting the analyses to the eight retained pairs"): break
    el = el.getnext()
moved = [copy.deepcopy(e) for e in els]
for e in els: body.remove(e)
body.remove(h_rc._p)
replace_markup(find(main, "The target result was mentioned more often for full-result"),
  "(4-year-olds: 5 responses, 3%; 5-year-olds: 6, 4%; adults: 2, 1%).",
  "(4-year-olds: 5 responses, 3%; 5-year-olds: 6, 4%; adults: 2, 1%). Among responses without an RVC, action-only responses predominated in all groups and in both result conditions (67%–90%), and children did not express the target result in separate predicates or without the action more often than adults (Supplemental Material, Section S11).")
main.save(D + "Manuscript_R18.docx"); print("saved main")

# ---------- 4. Supplemental Material ----------
sup = old
sb = sup.element.body
tblPr0 = copy.deepcopy(sup.tables[0]._tbl.tblPr)
for e in tblPr0.findall(qn("w:tblLayout")): tblPr0.remove(e)
title_txt = sup.paragraphs[0].text
def grab(start_text, end_test):
    start = [p for p in sup.paragraphs if p.text == start_text and p.style.name == "Table Number"][0]._p
    out = []; el = start
    while el is not None and el.tag != qn("w:sectPr"):
        out.append(copy.deepcopy(el))
        if end_test(el): break
        el = el.getnext()
    return out
isnote = lambda el: el.tag == qn("w:p") and Paragraph(el, None).text.startswith("Note.")
tabS1 = grab("Table S1", isnote)
sens = []
for p in sup.paragraphs:
    if p.text.startswith(("With all 10 pairs", "When each of the eight retained pairs")): sens.append(copy.deepcopy(p._p))
item_par = copy.deepcopy([p for p in sup.paragraphs if p.text.startswith("Table S2 shows")][0]._p)
tabItem = grab("Table S2", isnote)
tmpl_norm = copy.deepcopy([p for p in sup.paragraphs if p.style.name == "Normal" and p.runs][0]._p)
for el in list(sb):
    if el.tag != qn("w:sectPr"): sb.remove(el)

def append_el(el): sb.insert(len(sb) - 1, el); return el
def para(text, style="Normal", center=False, noindent=False):
    p = sup.add_paragraph(style=style)
    tmpl = copy.deepcopy(tmpl_norm.findall(qn("w:r"))[0])
    for e in markup_runs(tmpl, text): p._p.append(e)
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if noindent: p.paragraph_format.first_line_indent = 0
    return p
def heading(code, title): return para(f"{code.replace('Section ', '')}. {title}", style="Heading 1")
def table(number, title, header, rows, note, center_from=2, font_cols=None, widths=None):
    para(number, style="Table Number"); para("*" + title + "*", style="Table Title")
    tbl = sup.add_table(rows=1 + len(rows), cols=len(header))
    tbl._tbl.remove(tbl._tbl.tblPr); tbl._tbl.insert(0, copy.deepcopy(tblPr0))
    for i, row in enumerate([header] + rows):
        for j, txt in enumerate(row):
            cell = tbl.rows[i].cells[j]; p = cell.paragraphs[0]; p.style = sup.styles["Table Text"]
            for e in markup_runs(p.add_run("")._r, txt): p._p.append(e)
            if j >= center_from: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i == 0:
                tcPr = cell._tc.get_or_add_tcPr(); bd = tcPr.makeelement(qn("w:tcBorders"), {})
                bd.append(bd.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "4", qn("w:space"): "0", qn("w:color"): "000000"})); tcPr.append(bd)
    trPr = tbl.rows[0]._tr.get_or_add_trPr(); trPr.append(trPr.makeelement(qn("w:tblHeader"), {}))
    if widths:
        from docx.shared import Inches
        tbl.autofit = False
        for row in tbl.rows:
            for j, w in enumerate(widths): row.cells[j].width = Inches(w)
        grid = tbl._tbl.tblGrid
        for j, gc in enumerate(grid.findall(qn("w:gridCol"))): gc.set(qn("w:w"), str(int(widths[j] * 1440)))
    p = para("*Note.* " + note, style="Table Note")
    return tbl

para(title_txt, style="Title"); para("Supplemental Material", style="Heading 1")
# S1
heading(*T.S1[:2]); [para(t) for t in T.S1[2]]
# S2
heading(*T.S2[:2]); [para(t) for t in T.S2[2]]
for e in tabS1: append_el(e)
# S3, S4
for S in (T.S3, T.S4):
    heading(*S[:2]); [para(t) for t in S[2]]
# S5 test sentences
heading(*T.S5[:2]); [para(t) for t in T.S5[2]]
comp = pd.read_csv("/home/user/childrvc/analysis/comp.csv")
nv = pd.read_csv("/home/user/childrvc/analysis/norming_video.csv")
py = {"bo": "bāo", "xian": "xiān", "bai": "bāi", "si": "sī", "chui": "chuī", "ti": "tī", "qia": "qiā", "tu": "tú", "ca": "cā", "jia": "jiā",
      "nie": "niē", "pai": "pāi", "an": "àn", "che": "chě", "hui": "huī", "tui": "tuī", "jiao": "jiào", "nao": "náo", "chong": "chōng"}
rows = []
for k in range(1, 11):
    for assoc in ("stronger", "weaker"):
        v = nv[(nv.pair == k) & (nv.assoc == assoc)].iloc[0].v1; vv = "chuí" if (k == 10 and assoc == "weaker") else py[v]
        sv = comp[(comp.pair == k) & (comp.assoc == assoc) & (comp.form == "V1")].sentence.iloc[0]
        sr = comp[(comp.pair == k) & (comp.assoc == assoc) & (comp.form == "RVC")].sentence.iloc[0]
        rows.append([(str(k) + ("ᵃ" if k in (7, 8) else "")) if assoc == "stronger" else "", "*" + vv + "*", sv, sr])
table("Table S2", "Test Questions in Experiment 1", ["Pair", "V1", "V1 question", "RVC question"], rows,
      "In each pair, the first row is the stronger-association V1. ᵃ Excluded from the analyses of Experiment 1.", center_from=9, widths=[0.6, 1.0, 2.2, 2.2])
# S6 instructions
heading(*T.S6[:2]); para(T.S6_INTRO)
for kind, txt in T.S6[2]:
    para(txt, style="Heading 3" if kind == "H3" else "Normal")
# S7 models
heading("Section S7", "Model Specifications and Parameter Estimates"); [para(t) for t in T.S7_TEXT]
fx = pd.read_csv("/home/user/childrvc/analysis/out/fixef_all.csv")
lab = {"m1_8": "Model 1 (Experiment 1, RVC vs. V1)", "m2_8": "Model 2 (Experiment 1, V1 association)", "p_rvc": "Experiment 2: RVC use", "p_target": "Experiment 2: Mention of target result"}
tn = {"Intercept": "Intercept", "form1": "Form (RVC)", "result1": "Result (Full)", "age1": "Age (4-year-olds)", "age2": "Age (5-year-olds)", "assoc1": "Association (Stronger)"}
def term(t): return " × ".join(tn[x] for x in t.split(":"))
rows = []
for m in ("m1_8", "m2_8", "p_rvc", "p_target"):
    rows.append(["*" + lab[m] + "*", "", "", "", "", ""])
    for _, r in fx[fx.model == m].iterrows():
        rows.append(["   " + term(r.term), f"{r.est:.2f}", f"{r.se:.2f}", f"[{r.lo:.2f}, {r.hi:.2f}]", f"{r.rhat:.3f}", f"{r.ess:.0f}"])
table("Table S3", "Population-Level Estimates (Log-Odds) of the Main Models", ["Term", "Estimate", "*SE*", "95% CrI", "R̂", "Bulk ESS"], rows,
      "Experiment 1 models use the eight retained pairs; Experiment 2 models use all 10 pairs. Sum coding; the level in parentheses is coded +1 (for Age Group, the named group vs. the grand mean). Estimates are posterior means.", center_from=1, widths=[2.9, 0.75, 0.6, 1.25, 0.6, 0.8])
# S8 sensitivity
heading("Section S8", "Sensitivity Analyses for Experiment 1")
for e in sens: append_el(e)
# S9 item level
heading("Section S9", "Item-Level Acceptance in Experiment 1")
append_el(item_par)
for e in tabItem: append_el(e)
set_text(Paragraph(tabItem[0], sup.element.body), "Table S4")
# S10 coding
heading(*T.S10[:2]); [para(t) for t in T.S10[2]]
table("Table S5", "Expressions Coded as Mentioning the Target Result in Experiment 2", ["Pair", "Object", "Target result", "Expressions"],
      [[str(k), "*" + o.split(' ')[0] + "* " + " ".join(o.split(' ')[1:]), "*" + r.split(' ')[0] + "* " + " ".join(r.split(' ')[1:]), e] for k, o, r, e in T.RES_TABLE],
      "Expressions are matched in the Chinese transcripts after removing words that contain the same characters without expressing the result (e.g., 开始 \"begin,\" 离开 \"leave,\" 倒水 \"pour water\").", center_from=9, widths=[0.5, 1.6, 1.7, 2.7])
# S11 response categories
heading("Section S11", "Response Categories and Additional Analyses for Experiment 2")
for e in moved: append_el(e)
# renumber moved / copied items
for p in sup.paragraphs:
    if p.style.name == "Table Number" and p.text == "Table S2" and p._p is not None and "Item" in (p._p.getnext().xpath("string(.)") if p._p.getnext() is not None else ""):
        set_text(p, "Table S4")
    if p.style.name == "Table Number" and p.text == "Table 2": set_text(p, "Table S6")
    if p.text.startswith("Table S2 shows acceptance"): replace_markup(p, "Table S2 shows", "Table S4 shows")
    if p.text.startswith("Among responses without an RVC (Table 2)"): replace_markup(p, "(Table 2)", "(Table S6)")
    if "sensitivity analysis of Experiment 1 (Section S2)" in p.text: replace_markup(p, "(Section S2)", "(Section S8)")
for p in sup.paragraphs:
    if p.style.name == "Heading 3" and p.text == "Response Categories": delete(p)
sup.save(D + "Manuscript_R18_Supplemental_Material.docx"); print("saved supplement")
