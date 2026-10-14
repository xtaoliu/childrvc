import copy, re, sys
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_COLOR_INDEX, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
sys.path.insert(0, sys.path[0])
from content import SECTIONS, TABLE1, TABLE1_NOTE, NEW_REFS

SRC, OUT = sys.argv[1], sys.argv[2]
doc = Document(SRC)
FONT = "Times New Roman"

# ---------- styles ----------
def get_style(name, kind=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, kind)

def base(st, *, size=12, bold=False, italic=False, align=None, first=None, left=None, hanging=None,
         line=2.0, before=0, after=0, keep_next=False, page_break=False, outline=None, next_style=None):
    st.font.name = FONT; st.font.size = Pt(size); st.font.bold = bold; st.font.italic = italic
    st.font.color.rgb = RGBColor(0, 0, 0)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    pf = st.paragraph_format
    pf.line_spacing = line; pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.alignment = align; pf.keep_with_next = keep_next; pf.page_break_before = page_break
    pf.first_line_indent = first; pf.left_indent = left
    if hanging is not None:
        pf.left_indent = hanging; pf.first_line_indent = -hanging
    if outline is not None:
        ppr = st.element.get_or_add_pPr()
        ol = ppr.find(qn("w:outlineLvl"))
        if ol is None:
            ol = OxmlElement("w:outlineLvl"); ppr.append(ol)
        ol.set(qn("w:val"), str(outline))
    if next_style is not None:
        st.next_paragraph_style = next_style

normal = doc.styles["Normal"]
base(normal, first=Inches(0.5))
# APA: no widow/orphan requirement; keep Word default
st_title = get_style("Title"); st_title.base_style = normal
base(st_title, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=Inches(0), next_style=normal)
h1 = get_style("Heading 1"); h1.base_style = normal
base(h1, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=Inches(0), keep_next=True, outline=0, next_style=normal)
h2 = get_style("Heading 2"); h2.base_style = normal
base(h2, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, first=Inches(0), keep_next=True, outline=1, next_style=normal)
h3 = get_style("Heading 3"); h3.base_style = normal
base(h3, bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, first=Inches(0), keep_next=True, outline=2, next_style=normal)
for h in (h1, h2, h3, st_title):
    h.font.color.rgb = RGBColor(0, 0, 0)

ex = get_style("Example"); ex.base_style = normal
base(ex, first=Inches(0), line=1.0, keep_next=True)
ex_tabs = ex.paragraph_format.tab_stops
ex_tabs.add_tab_stop(Inches(0.5)); ex_tabs.add_tab_stop(Inches(0.875))
ex_last = get_style("Example Last"); ex_last.base_style = ex
base(ex_last, first=Inches(0), line=1.0, after=12)
ex_last.paragraph_format.tab_stops.add_tab_stop(Inches(0.5)); ex_last.paragraph_format.tab_stops.add_tab_stop(Inches(0.875))
ex_first = get_style("Example First"); ex_first.base_style = ex
base(ex_first, first=Inches(0), line=1.0, before=12, keep_next=True)
ex_first.paragraph_format.tab_stops.add_tab_stop(Inches(0.5)); ex_first.paragraph_format.tab_stops.add_tab_stop(Inches(0.875))

rq = get_style("Research Question"); rq.base_style = normal
base(rq, first=Inches(0), left=Inches(0.5))
ref = get_style("Reference"); ref.base_style = normal
base(ref, hanging=Inches(0.5))
tnum = get_style("Table Number"); tnum.base_style = normal
base(tnum, bold=True, first=Inches(0), keep_next=True, before=12)
ttitle = get_style("Table Title"); ttitle.base_style = normal
base(ttitle, italic=True, first=Inches(0), keep_next=True)
ttext = get_style("Table Text"); ttext.base_style = normal
base(ttext, size=10, first=Inches(0), line=1.0)
tnote = get_style("Table Note"); tnote.base_style = normal
base(tnote, size=10, first=Inches(0), line=1.0, before=6, after=12)
fn = get_style("Footnote Text")
base(fn, size=10, first=Inches(0), line=1.0)

# ---------- helpers ----------
KEEP_RPR = {qn("w:i"), qn("w:iCs"), qn("w:highlight"), qn("w:vertAlign"), qn("w:rStyle"), qn("w:u")}

def clean_paragraph(p, style, strip_bold=False):
    ppr = p._p.get_or_add_pPr()
    for child in list(ppr):
        if child.tag not in (qn("w:pStyle"), qn("w:sectPr"), qn("w:rPr")):
            ppr.remove(child)
    p.style = style
    for r in p._p.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is None: continue
        for c in list(rpr):
            if c.tag in KEEP_RPR: continue
            if c.tag in (qn("w:b"), qn("w:bCs")) and not strip_bold: continue
            rpr.remove(c)

TOKEN = re.compile(r"(\*[^*]+\*|⟦[^⟧]+⟧)")
def add_runs(p, text, italic_all=False):
    text = re.sub(r'"([^"]*)"', "\u201C\\1\u201D", text)
    text = re.sub(r"(\w)'", "\\1\u2019", text)
    text = text.replace("2^10", "2\u00b9\u2070")
    for seg in TOKEN.split(text):
        if not seg: continue
        if seg.startswith("*") and seg.endswith("*"):
            r = p.add_run(seg[1:-1]); r.italic = not italic_all
        elif seg.startswith("⟦"):
            r = p.add_run(seg); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            r = p.add_run(seg)
            if italic_all: r.italic = True
    return p

body = doc.element.body
paras = list(doc.paragraphs)

# ---------- restyle the Introduction (author's version) ----------
ref_head_idx = next(i for i, p in enumerate(paras) if p.text.strip() == "References")
heading2_texts = {"V1–V2 Combination in Mandarin RVCs",
                  "Monomorphemic V1s: Classification, Result Likelihood, and Action Typicality",
                  "Children’s Comprehension of RVCs and V1s", "Children’s Production of RVCs and V1s",
                  "The Present Study"}
example_idx = []
for i, p in enumerate(paras[:ref_head_idx]):
    t = p.text
    if i == 0:
        clean_paragraph(p, st_title, strip_bold=True)
    elif t.strip() in heading2_texts:
        clean_paragraph(p, h2, strip_bold=True)
    elif re.match(r"^\(\d\)\t|^\t", t):
        example_idx.append(i); clean_paragraph(p, ex)
    elif re.match(r"^RQ\d[.:]", t):
        clean_paragraph(p, rq)
    else:
        clean_paragraph(p, normal)
# first/last line of each example block get spacing
for i in example_idx:
    if i - 1 not in example_idx: paras[i].style = ex_first
    if i + 1 not in example_idx: paras[i].style = ex_last

# small fixes in the author's text
def replace_in_par(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new); return True
    runs = p.runs; texts = [r.text for r in runs]; full = "".join(texts)
    k = full.find(old)
    if k < 0: return False
    end = k + len(old); pos = 0; first = True
    for r, t in zip(runs, texts):
        a, b = pos, pos + len(t); pos = b
        if b <= k or a >= end: continue
        if r.italic: raise SystemExit("span over italic: " + old)
        lo, hi = max(k, a) - a, min(end, b) - a
        r.text = t[:lo] + (new if first else "") + t[hi:]
        first = False
    return True
EVENT_FIXES = [
 ("depend on which V1 is used and on how that action is related", "depend on which V1 is used and on how the action it names is related"),
 ("pairs of V1s that take the same object and share a target result", "pairs of V1s that take the same object and combine with the same V2"),
 ("what makes a given V1 a ", "what makes the action named by a given V1 a "),
 ("a V1’s relation to a result depends on the result in question", "the relation between the action named by a V1 and a result depends on the result in question"),
 ("it characterizes the selected V1–object expressions by their association", "it characterizes the actions named by the selected V1–object expressions by their association"),
 ("The relation between a V1 action and a target result", "The relation between the action named by a V1 and a target result"),
 ("Once the V1–object action has been performed", "Once the action named by the V1–object expression has been performed"),
 ("how a particular V1 event relates to a particular result", "how a particular action relates to a particular result"),
 ("A V1–object expression has a stronger association with the target result than its partner when", "The action named by a V1–object expression has a stronger association with the target result than the action named by its partner when"),
 ("adult reference measures of the selected V1–object–result relations", "adult reference measures of the relations between the selected actions and their target results"),
 ("comparing V1s that act on the same object and share a target result", "comparing V1s that name actions on the same object with the same target result"),
 ("so that one V1–object expression was expected to be more strongly associated", "so that one action was expected to be more strongly associated"),
]
applied = set()
for p in paras[:ref_head_idx]:
    if "Section 1.2 addresses this relation.”" in p.text:
        replace_in_par(p, "Section 1.2 addresses this relation.”", "The next section addresses this relation.")
    if p.text.startswith("RQ3:"):
        replace_in_par(p, "RQ3:", "RQ3.")
    for old, new in EVENT_FIXES:
        if old in p.text:
            replace_in_par(p, old, new); applied.add(old)
missing = [o for o, n in EVENT_FIXES if o not in applied]
if missing: raise SystemExit("not applied: " + repr(missing))

# footnote text style
try:
    fpart = doc.part._footnotes_part if hasattr(doc.part, "_footnotes_part") else None
except Exception:
    fpart = None

# ---------- References ----------
ref_head = paras[ref_head_idx]
clean_paragraph(ref_head, h1, strip_bold=True)
ref_head.paragraph_format.page_break_before = True
old_refs = paras[ref_head_idx + 1:]
for p in old_refs:
    clean_paragraph(p, ref)

def make_par(style):
    p = doc.add_paragraph(style=style)
    return p

def insert_before(new_p, anchor_p):
    anchor_p._p.addprevious(new_p._p)

# new refs, sorted insertion
def key(s): return re.sub(r"[^A-Za-zÀ-ÿ ]", "", s).lower()
for s in NEW_REFS:
    np_ = make_par(ref)
    # italics in refs: *...*
    add_runs(np_, s)
    allp = doc.paragraphs
    start = next(i for i, p in enumerate(allp) if p._p is ref_head._p)
    anchor = next((p for p in allp[start + 1:] if p.style.name == "Reference" and p._p is not np_._p and key(p.text) > key(s)), None)
    if anchor is not None:
        insert_before(np_, anchor)

# ---------- new sections, inserted before References heading ----------
style_for = {"H1": h1, "H2": h2, "H3": h3, "P": normal, "Q": normal}

def build_table1(anchor):
    cap1 = make_par(tnum); add_runs(cap1, "Table 1"); insert_before(cap1, anchor)
    cap2 = make_par(ttitle); add_runs(cap2, "Action Pairs, Adult Ratings, and Corpus Frequency of the Corresponding RVCs", italic_all=True); insert_before(cap2, anchor)
    hdr = ["Pair", "Object", "Target result (V2)", "V1", "RVC", "Likelihood\n*M* (*SD*)", "Typicality\n*M* (*SD*)", "CCL\ndocuments"]
    tbl = doc.add_table(rows=1, cols=len(hdr))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    def cell_write(cell, text, align=WD_ALIGN_PARAGRAPH.LEFT):
        cell.text = ""
        lines = text.split("\n")
        p = cell.paragraphs[0]; p.style = ttext; p.alignment = align
        for k, ln in enumerate(lines):
            if k: p.add_run().add_break()
            add_runs(p, ln)
    for j, h in enumerate(hdr):
        cell_write(tbl.rows[0].cells[j], h, WD_ALIGN_PARAGRAPH.CENTER if j >= 5 else WD_ALIGN_PARAGRAPH.LEFT)
    for pair, obj, objg, v2, v2g, rows, excl in TABLE1:
        for k, (v1, v1g, L, T, ccl) in enumerate(rows):
            r = tbl.add_row().cells
            if k == 0:
                cell_write(r[0], pair); cell_write(r[1], f"*{obj}*\n\"{objg}\""); cell_write(r[2], f"*{v2}*\n\"{v2g}\"")
            cell_write(r[3], f"*{v1}*\n\"{v1g}\""); cell_write(r[4], f"*{v1}-{v2}*")
            cell_write(r[5], L, WD_ALIGN_PARAGRAPH.CENTER); cell_write(r[6], T, WD_ALIGN_PARAGRAPH.CENTER); cell_write(r[7], ccl, WD_ALIGN_PARAGRAPH.CENTER)
    # APA borders: top & bottom of table, below header; pair separators none
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:color"), "000000"); borders.append(e)
    for edge in ("left", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "nil"); borders.append(e)
    tblPr.append(borders)
    for c in tbl.rows[0].cells:
        tcPr = c._tc.get_or_add_tcPr(); tb = OxmlElement("w:tcBorders")
        b = OxmlElement("w:bottom"); b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "8"); b.set(qn("w:color"), "000000"); tb.append(b); tcPr.append(tb)
    # thin rule after each pair (after 2nd row of each pair) for readability: skip (APA minimal rules)
    widths = [0.35, 0.8, 0.95, 0.75, 0.95, 0.95, 0.9, 0.85]
    tbl.autofit = False
    tblW = OxmlElement("w:tblW"); tblW.set(qn("w:w"), str(int(sum(widths) * 1440))); tblW.set(qn("w:type"), "dxa")
    old = tblPr.find(qn("w:tblW"))
    if old is not None: tblPr.remove(old)
    tblPr.append(tblW)
    mar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), "58"); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for row in tbl.rows:
        for j, c in enumerate(row.cells):
            c.width = Inches(widths[j])
    grid = tbl._tbl.find(qn("w:tblGrid"))
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[j] * 1440)))
    # header row repeat
    trPr = tbl.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    # merge pair/object/result cells vertically
    for i in range(len(TABLE1)):
        a, b = tbl.rows[1 + 2 * i], tbl.rows[2 + 2 * i]
        for j in range(3):
            a.cells[j].merge(b.cells[j])
    insert_before(tbl, anchor) if False else anchor._p.addprevious(tbl._tbl)
    note = make_par(tnote)
    r = note.add_run("Note."); r.italic = True
    add_runs(note, " " + TABLE1_NOTE)
    insert_before(note, anchor)

for kind, text in SECTIONS:
    if kind == "TABLE1":
        build_table1(ref_head); continue
    if kind == "EX":
        lines = text
        for k, (num, let, ln) in enumerate(lines):
            st = ex_first if k == 0 else (ex_last if k == len(lines) - 1 else ex)
            p = make_par(st)
            add_runs(p, f"{num}\t{let}\t{ln}" if (num or let) else f"\t\t{ln}")
            insert_before(p, ref_head)
        continue
    if kind == "P4":
        head, rest = text
        p = make_par(normal)
        r = p.add_run(head); r.bold = True
        add_runs(p, " " + rest)
        insert_before(p, ref_head); continue
    p = make_par(style_for[kind])
    add_runs(p, text)
    insert_before(p, ref_head)

doc.save(OUT)
print("saved", OUT)
