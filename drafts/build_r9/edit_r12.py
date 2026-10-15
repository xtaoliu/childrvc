import re, sys, copy
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.paragraph import Paragraph
SRC, OUT = sys.argv[1], sys.argv[2]
doc = Document(SRC)
P = doc.paragraphs

def find(prefix):
    hits = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(hits) == 1, (prefix, len(hits)); return hits[0]

def replace_span(p, old, new):
    runs = p.runs; texts = [r.text for r in runs]; full = "".join(texts)
    k = full.find(old); assert k >= 0, old
    end = k + len(old); pos = 0; first = True
    for r, t in zip(runs, texts):
        a, b = pos, pos + len(t); pos = b
        if b <= k or a >= end: continue
        assert not r.italic, "italic in span: " + old
        lo, hi = max(k, a) - a, min(end, b) - a
        r.text = t[:lo] + (new if first else "") + t[hi:]; first = False

TOKEN = re.compile(r"(\*[^*]+\*|⟦[^⟧]+⟧)")
def add_runs(p, text):
    text = re.sub(r'"([^"]*)"', "\u201C\\1\u201D", text)
    text = re.sub(r"(\w)'", "\\1\u2019", text)
    for seg in TOKEN.split(text):
        if not seg: continue
        if seg.startswith("*") and seg.endswith("*"):
            r = p.add_run(seg[1:-1].replace("{AST}", "*")); r.italic = True
        elif seg.startswith("⟦"):
            r = p.add_run(seg); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            p.add_run(seg.replace("{AST}", "*"))

def clear(p):
    for r in list(p._p):
        if r.tag.endswith("}r") or r.tag.endswith("}hyperlink"): p._p.remove(r)

def insert_after(p, text):
    new = copy.deepcopy(p._p)
    for r in list(new):
        if not r.tag.endswith("}pPr"): new.remove(r)
    p._p.addnext(new)
    np_ = Paragraph(new, p._parent); add_runs(np_, text); return np_

# ---- 3. Strong/Weak ----
p47 = find("The materials comprise 10 pairs of actions")
replace_span(p47, " The V1s were originally selected as Strong and Weak V1s with the preverbal time-span adverbial diagnostic (Martin et al., 2020; Martin et al., 2023). In the present study, the two sets are characterized by the adult ratings described below rather than by this diagnostic.", "")
insert_after(p47, "When the materials were designed, the two sets were defined as Strong and Weak V1s, following Martin et al. (2020, 2023). On this classification, Strong V1s, such as *xǐ* \"wash,\" strongly imply a change of state in the object, whereas Weak V1s, such as *mō* \"pet,\" at best weakly imply one. The two classes are distinguished by a single diagnostic: Only Strong V1s are compatible with a preverbal time-span adverbial, as in *yì fēnzhōng xǐ-le yì shuāng wàzi* \"washed a pair of socks in one minute\" versus *{AST}yì fēnzhōng mō-le yì zhī māo* \"petted a cat in one minute.\" Some V1s in the present materials were taken from Martin et al. (2020), and the others were selected with this diagnostic. We no longer use this classification, for three reasons. First, the diagnostic is binary, whereas the property it was meant to capture, the strength with which a V1 implies a result, is a matter of degree. Second, the diagnostic tracks whether a result construal of the verb phrase is available, which the object and the context can supply, rather than a fixed property of the V1: Several V1s in the weaker set are also compatible with the adverbial when combined with a suitable object, as in *yì fēnzhōng tuī-le yí ge xiāngzi* \"pushed a box in one minute,\" and the same holds for *chōng* \"rinse\" and *chuí* \"hammer.\" Third, the stronger set combined verbs that name an action directed at a specific result, such as *bāo* \"peel,\" *sī* \"tear,\" and *cā* \"wipe,\" with verbs that name a manner of action compatible with many results, such as *tī* \"kick\" and *chuī* \"blow,\" and the diagnostic provides no independent reason to group them together. Because the diagnostic characterizes verbs rather than the relation between an action and a particular result, the present study characterizes the pairs by adult ratings of each action's association with its target result, described below. The labels *stronger-association* and *weaker-association* refer to these ratings, not to the original classification.")

# ---- 1 & 2. ratings: instructions, scale labels, list order ----
p55 = find("Twenty native speakers of Mandarin (21–35 years")
clear(p55)
add_runs(p55, "Twenty native speakers of Mandarin (21–35 years, *M* = 28.2) ⟦性别与招募方式待补⟧ completed an online questionnaire of about 5 min, in which they rated each of the 20 V1–object expressions in two separate blocks. After a consent statement and questions about their language background (native language, frequency of Mandarin use, and English proficiency), they were told that the questions had no correct answers and were asked to respond on the basis of their everyday experience and first intuition. In the likelihood block, each item presented a perfective sentence containing the V1–object expression, followed by a question about the target result, for example, *Xiǎo Lǐ sī-le zhǐ* \"Xiao Li tore the paper. How likely do you think it is that the paper ended up torn?\" The instructions asked participants to judge, on the basis of the sentence and their everyday experience, how likely it was that the stated result eventually occurred. They typed a number from 0 to 100, where 0 was labeled *cannot happen* and 100 *will certainly happen*. In the typicality block, the instructions stated that each item first named a result that someone wanted to achieve and then an action that could be used to achieve it, and asked participants to rate how typical the action was as a way of achieving that result, for example, \"Suppose Xiao Li wants to break the walnut. How typical is *jiā hétao* ‘clamping the walnut’ as a way of doing this?\" Responses were given on a 7-point scale whose endpoints were labeled 1 (*very atypical*) and 7 (*very typical*); the intermediate points were numbered but not labeled. The likelihood items thus presented each expression in the same perfective form as the V1 questions of Experiment 1. Because the ratings were collected for sentences naming the expressions, they characterize the action–result relation that each expression describes, not the particular videos.")
insert_after(p55, "Four lists were constructed, and five raters completed each list. In all lists, the 20 likelihood items preceded the 20 typicality items. Within each block, the 10 expressions of one set were presented first and the 10 expressions of the other set second, with the pairs in the same order in both halves, so that the two members of each pair were always separated by nine other items. The stronger-association set came first in Lists 1 and 3, and the weaker-association set came first in Lists 2 and 4. Four different pair orders were used, two for Lists 1 and 2 and two for Lists 3 and 4, and the two blocks of each list used different orders. Because the likelihood block always came first, the typicality judgments may have been influenced by the preceding likelihood judgments. All raters completed all items, and no ratings were excluded. The ratings were collected after the child data had been collected and initially analyzed with the original Strong/Weak coding; the raters did not take part in the child experiments and received no information about children's responses.")

# ---- 4. why conventionality ----
p61 = find("To describe how conventional each RVC is")
replace_span(p61, "To describe how conventional each RVC is, we searched", "We also examined how conventional each RVC is, for two reasons. First, as noted in the Introduction, V1–V2 combinations differ in how conventional they are, and an action that reliably brings about a result is likely to be described often with the corresponding RVC; corpus frequency thus characterizes the pairs independently of the adult ratings. Second, children may be more familiar with the RVCs formed with stronger-association V1s, and familiarity with an RVC could itself affect how they interpret the corresponding V1 (see General Discussion). We therefore searched")

# ---- limitation: fixed block order ----
p119 = find("Several limitations follow from the design.")
replace_span(p119, "with the two ratings collected from different raters.", "with the two ratings collected from different raters; in the present study, the same raters always gave the likelihood ratings first.")

doc.save(OUT); print("ok")
